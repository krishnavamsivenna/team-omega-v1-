import io
import re
from typing import Dict, Any, List
import pypdf
import docx
from app.services.skill_taxonomy import extract_skills_from_text

class ResumeParser:
    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
        """Extracts text from PDF, DOCX, or TXT file bytes."""
        extension = filename.lower().split(".")[-1]
        
        if extension == "pdf":
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            return "\n".join(pages_text).strip()
            
        elif extension in ["docx", "doc"]:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs).strip()
            
        elif extension in ["txt", "md"]:
            try:
                return file_bytes.decode("utf-8").strip()
            except UnicodeDecodeError:
                return file_bytes.decode("latin-1", errors="ignore").strip()
        else:
            raise ValueError(f"Unsupported file format: .{extension}. Supported formats are PDF, DOCX, TXT.")

    @staticmethod
    def parse_contact_info(text: str) -> Dict[str, Any]:
        """Extracts contact details such as email, phone, and links."""
        # Email
        email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
        emails = re.findall(email_pattern, text)
        email = emails[0] if emails else None
        
        # Phone: matches various US and international formats
        phone_pattern = r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
        phones = re.findall(phone_pattern, text)
        phone = phones[0] if phones else None
        
        # Links: GitHub, LinkedIn, portfolios
        links_pattern = r'(?:https?:\/\/)?(?:www\.)?(?:github\.com\/[a-zA-Z0-9_-]+|linkedin\.com\/in\/[a-zA-Z0-9_-]+|[a-zA-Z0-9-]+\.(?:io|dev|com|org|net)(?:\/[^\s]*)?)'
        raw_links = re.findall(links_pattern, text, re.IGNORECASE)
        # Filter out common false positives like email domains
        links = []
        for l in raw_links:
            if email and email.endswith(l):
                continue
            if not l.startswith("http"):
                l = "https://" + l
            if l not in links:
                links.append(l)

        # Basic name heuristic (first line or first non-empty lines before email)
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        name = None
        for line in lines[:5]:
            if len(line.split()) in [2, 3, 4] and not re.search(r'[@\(\)\d]', line):
                name = line
                break

        return {
            "name": name,
            "email": email,
            "phone": phone,
            "links": links[:5]
        }

    @staticmethod
    def detect_sections(text: str) -> List[str]:
        """Detects standard resume sections present in the text."""
        standard_sections = {
            "experience": r"(work\s+experience|professional\s+experience|employment\s+history|experience)",
            "education": r"(education|academic\s+background|qualifications|degrees)",
            "skills": r"(skills|technical\s+skills|core\s+competencies|technologies)",
            "projects": r"(projects|personal\s+projects|open\s+source)",
            "certifications": r"(certifications|licenses|courses|awards)",
            "summary": r"(summary|professional\s+summary|objective|about\s+me)"
        }
        
        detected = []
        for section, pattern in standard_sections.items():
            if re.search(r"(?i)(?:^|\n)\s*" + pattern + r"\s*(?::|\n|$)", text):
                detected.append(section)
        return detected

    @classmethod
    def parse_resume(cls, text: str) -> Dict[str, Any]:
        """Parses full structured resume data from raw text."""
        contact = cls.parse_contact_info(text)
        detected_sections = cls.detect_sections(text)
        skills, categorized_skills = extract_skills_from_text(text)
        words = text.split()
        
        return {
            "contact": contact,
            "detected_sections": detected_sections,
            "skills": skills,
            "categorized_skills": categorized_skills,
            "word_count": len(words),
        }
